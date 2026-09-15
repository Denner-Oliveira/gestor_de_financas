from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unicodedata import combining, normalize
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from backend.app.core.security import obter_usuario_atual
from backend.app.db.dependencies import obter_sessao
from backend.app.models.models import Compra, ContaFinanceira, Receita, Usuario
from backend.app.routers.lancamentos import _criar_parcelas

router = APIRouter(tags=['financeiro'])

IMPORT_HEADERS = (
    "tipo",
    "descricao",
    "categoria",
    "valor",
    "data",
    "conta",
    "quantidade_parcelas",
    "observacoes",
)

def _ler_data_importada(valor: object, linha: int) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        try:
            return date.fromisoformat(valor.strip())
        except ValueError:
            pass
    raise HTTPException(
        status_code=422,
        detail=f"Linha {linha}: data inválida. Use o formato AAAA-MM-DD.",
    )


def _texto_importado(valor: object, campo: str, linha: int) -> str:
    texto = str(valor).strip() if valor is not None else ""
    if not texto:
        raise HTTPException(status_code=422, detail=f"Linha {linha}: {campo} é obrigatório.")
    return texto


def _normalizar_tipo_importado(valor: object, linha: int) -> str:
    texto = _texto_importado(valor, "tipo", linha)
    normalizado = normalize("NFKD", texto)
    normalizado = "".join(caractere for caractere in normalizado if not combining(caractere))
    normalizado = " ".join(normalizado.casefold().split())
    tipos = {
        "receita": "receita",
        "receitas": "receita",
        "entrada": "receita",
        "entradas": "receita",
        "compra": "compra",
        "compras": "compra",
        "despesa": "compra",
        "despesas": "compra",
        "saida": "compra",
        "saidas": "compra",
    }
    tipo = tipos.get(normalizado)
    if tipo is None:
        raise HTTPException(
            status_code=422,
            detail=(
                f"Linha {linha}: tipo inválido ({texto!r}). "
                "Use receita ou compra."
            ),
        )
    return tipo



@router.get("/importacoes/modelo")
def baixar_modelo_importacao(
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
) -> StreamingResponse:
    conta = sessao.scalar(
        select(ContaFinanceira)
        .where(
            ContaFinanceira.usuario_id == usuario.id,
            ContaFinanceira.ativa.is_(True),
        )
        .order_by(ContaFinanceira.id)
    )
    nome_conta = conta.nome if conta is not None else "Nome da conta cadastrada"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Lançamentos"
    worksheet.append(list(IMPORT_HEADERS))
    worksheet.append(
        [
            "receita",
            "Salário",
            "Salário",
            5000,
            date.today(),
            nome_conta,
            1,
            "Exemplo de receita",
        ]
    )
    worksheet.append(
        [
            "compra",
            "Supermercado",
            "Alimentação",
            250.50,
            date.today(),
            nome_conta,
            1,
            "Exemplo de compra",
        ]
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = cell.font.copy(bold=True)
    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = 22
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)
    return StreamingResponse(
        arquivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo-lancamentos.xlsx"'},
    )


@router.post("/importacoes/lancamentos")
def importar_lancamentos(
    arquivo: UploadFile = File(...),
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    nome = (arquivo.filename or "").lower()
    if not nome.endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="Envie um arquivo XLSX.")
    conteudo = arquivo.file.read()
    if len(conteudo) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="O arquivo deve ter no máximo 10 MB.")
    try:
        workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except (TypeError, ValueError, OSError) as exc:
        raise HTTPException(status_code=422, detail="Arquivo XLSX inválido.") from exc

    worksheet = workbook.active
    linhas = worksheet.iter_rows(values_only=True)
    cabecalho = tuple(
        str(valor).strip().lower() if valor is not None else ""
        for valor in next(linhas, ())
    )
    if cabecalho != IMPORT_HEADERS:
        raise HTTPException(
            status_code=422,
            detail=f"Cabeçalho inválido. Use o modelo com as colunas: {', '.join(IMPORT_HEADERS)}.",
        )

    contas = {
        conta.nome.casefold(): conta
        for conta in sessao.scalars(
            select(ContaFinanceira).where(
                ContaFinanceira.usuario_id == usuario.id,
                ContaFinanceira.ativa.is_(True),
            )
        )
    }
    pendentes: list[tuple[str, dict[str, object]]] = []
    for numero_linha, valores in enumerate(linhas, start=2):
        if not any(valor is not None and str(valor).strip() for valor in valores):
            continue
        dados = dict(zip(IMPORT_HEADERS, valores, strict=False))
        tipo = _normalizar_tipo_importado(dados["tipo"], numero_linha)
        nome_conta = _texto_importado(dados["conta"], "conta", numero_linha)
        if nome_conta.casefold() not in contas:
            raise HTTPException(
                status_code=422,
                detail=f"Linha {numero_linha}: conta '{nome_conta}' não encontrada.",
            )
        try:
            valor = Decimal(str(dados["valor"])).quantize(Decimal("0.01"))
        except (ArithmeticError, ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: valor inválido.") from exc
        if valor <= 0:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: valor deve ser maior que zero.")
        quantidade = dados["quantidade_parcelas"] or 1
        try:
            quantidade = int(quantidade)
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: quantidade_parcelas inválida.") from exc
        if tipo == "receita" and quantidade != 1:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: receitas devem ter 1 parcela.")
        if not 1 <= quantidade <= 120:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: quantidade_parcelas deve estar entre 1 e 120.")
        pendentes.append(
            (
                tipo,
                {
                    "descricao": _texto_importado(dados["descricao"], "descricao", numero_linha),
                    "categoria": _texto_importado(dados["categoria"], "categoria", numero_linha),
                    "valor": valor,
                    "data": _ler_data_importada(dados["data"], numero_linha),
                    "conta_id": contas[nome_conta.casefold()].id,
                    "quantidade_parcelas": quantidade,
                    "observacoes": str(dados["observacoes"]).strip() if dados["observacoes"] is not None else None,
                },
            )
        )

    for tipo, dados in pendentes:
        if tipo == "receita":
            sessao.add(
                Receita(
                    usuario_id=usuario.id,
                    descricao=dados["descricao"],
                    categoria=dados["categoria"],
                    valor=dados["valor"],
                    data=dados["data"],
                    conta_id=dados["conta_id"],
                    observacoes=dados["observacoes"],
                )
            )
        else:
            compra = Compra(
                usuario_id=usuario.id,
                descricao=dados["descricao"],
                categoria=dados["categoria"],
                valor_total=dados["valor"],
                quantidade_parcelas=dados["quantidade_parcelas"],
                data_compra=dados["data"],
                conta_id=dados["conta_id"],
                observacoes=dados["observacoes"],
            )
            _criar_parcelas(compra)
            sessao.add(compra)
    sessao.commit()
    return {"importados": len(pendentes)}
