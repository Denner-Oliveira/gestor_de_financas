from datetime import date, datetime
from decimal import Decimal, ROUND_DOWN
from io import BytesIO
from unicodedata import combining, normalize

from openpyxl import Workbook, load_workbook
from dateutil.relativedelta import relativedelta
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from backend.app.core.security import obter_usuario_atual
from backend.app.db.dependencies import obter_sessao
from backend.app.models.models import (
    Compra,
    ContaFinanceira,
    Parcela,
    Receita,
    Usuario,
)
from backend.app.schemas.schemas import (
    ContaFinanceiraCriar,
    ContaFinanceiraAtualizar,
    ContaFinanceiraResposta,
    CompraCriar,
    CompraAtualizar,
    CompraResposta,
    ReceitaAtualizar,
    ReceitaCriar,
    ReceitaResposta,
)


router = APIRouter(tags=["financeiro"])

IMPORT_HEADERS = (
    "tipo",
    "descricao",
    "categoria",
    "valor",
    "data",
    "conta",
    "quantidade_parcelas",
    "observacoes",
)


def _ler_data_importada(valor: object, linha: int) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        try:
            return date.fromisoformat(valor.strip())
        except ValueError:
            pass
    raise HTTPException(
        status_code=422,
        detail=f"Linha {linha}: data inválida. Use o formato AAAA-MM-DD.",
    )


def _texto_importado(valor: object, campo: str, linha: int) -> str:
    texto = str(valor).strip() if valor is not None else ""
    if not texto:
        raise HTTPException(status_code=422, detail=f"Linha {linha}: {campo} é obrigatório.")
    return texto


def _normalizar_tipo_importado(valor: object, linha: int) -> str:
    texto = _texto_importado(valor, "tipo", linha)
    normalizado = normalize("NFKD", texto)
    normalizado = "".join(caractere for caractere in normalizado if not combining(caractere))
    normalizado = " ".join(normalizado.casefold().split())
    tipos = {
        "receita": "receita",
        "receitas": "receita",
        "entrada": "receita",
        "entradas": "receita",
        "compra": "compra",
        "compras": "compra",
        "despesa": "compra",
        "despesas": "compra",
        "saida": "compra",
        "saidas": "compra",
    }
    tipo = tipos.get(normalizado)
    if tipo is None:
        raise HTTPException(
            status_code=422,
            detail=(
                f"Linha {linha}: tipo inválido ({texto!r}). "
                "Use receita ou compra."
            ),
        )
    return tipo


def _criar_parcelas(compra: Compra) -> None:
    valor_parcela = (compra.valor_total / compra.quantidade_parcelas).quantize(
        Decimal("0.01"), rounding=ROUND_DOWN
    )
    valores = [valor_parcela] * compra.quantidade_parcelas
    valores[-1] += compra.valor_total - sum(valores)
    compra.parcelas = [
        Parcela(
            numero=numero,
            valor=valor,
            data_vencimento=compra.data_compra + relativedelta(months=numero - 1),
        )
        for numero, valor in enumerate(valores, start=1)
    ]


def _validar_conta(
    conta_id: int,
    usuario_id: int,
    sessao: Session,
) -> ContaFinanceira:
    conta = sessao.scalar(
        select(ContaFinanceira).where(
            ContaFinanceira.id == conta_id,
            ContaFinanceira.usuario_id == usuario_id,
            ContaFinanceira.ativa.is_(True),
        )
    )
    if conta is None:
        raise HTTPException(status_code=404, detail="Conta financeira não encontrada.")
    return conta


@router.post(
    "/contas",
    response_model=ContaFinanceiraResposta,
    status_code=status.HTTP_201_CREATED,
)
def criar_conta(
    dados: ContaFinanceiraCriar,
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    conta = ContaFinanceira(**dados.model_dump(), usuario_id=usuario.id)
    sessao.add(conta)
    sessao.commit()
    sessao.refresh(conta)
    return conta


@router.get("/contas", response_model=list[ContaFinanceiraResposta])
def listar_contas(
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    return list(
        sessao.scalars(
            select(ContaFinanceira)
            .where(
                ContaFinanceira.usuario_id == usuario.id,
                ContaFinanceira.ativa.is_(True),
            )
            .order_by(ContaFinanceira.banco, ContaFinanceira.nome)
        )
    )


@router.get("/relatorios/geral")
def relatorio_geral(
    inicio: date,
    fim: date,
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    if inicio > fim:
        raise HTTPException(
            status_code=422,
            detail="A data inicial deve ser anterior ou igual à data final.",
        )

    receitas = list(
        sessao.scalars(
            select(Receita).where(
                Receita.usuario_id == usuario.id,
                Receita.data >= inicio,
                Receita.data <= fim,
            )
        )
    )
    compras = list(
        sessao.scalars(
            select(Compra)
            .options(selectinload(Compra.parcelas))
            .join(Compra.parcelas)
            .where(
                Compra.usuario_id == usuario.id,
                Parcela.data_vencimento >= inicio,
                Parcela.data_vencimento <= fim,
            )
        ).unique()
    )
    contas = {
        conta.id: conta.nome
        for conta in sessao.scalars(
            select(ContaFinanceira).where(
                ContaFinanceira.usuario_id == usuario.id,
            )
        )
    }

    por_conta: dict[int, dict[str, object]] = {
        conta_id: {"conta": nome, "receitas": Decimal("0"), "despesas": Decimal("0")}
        for conta_id, nome in contas.items()
    }
    por_categoria: dict[str, Decimal] = {}
    por_mes: dict[str, dict[str, Decimal]] = {}

    for receita in receitas:
        conta = por_conta.setdefault(
            receita.conta_id or 0,
            {"conta": "Conta não identificada", "receitas": Decimal("0"), "despesas": Decimal("0")},
        )
        conta["receitas"] += receita.valor
        mes = receita.data.strftime("%Y-%m")
        resumo_mes = por_mes.setdefault(
            mes, {"receitas": Decimal("0"), "despesas": Decimal("0")}
        )
        resumo_mes["receitas"] += receita.valor

    for compra in compras:
        conta = por_conta.setdefault(
            compra.conta_id or 0,
            {"conta": "Conta não identificada", "receitas": Decimal("0"), "despesas": Decimal("0")},
        )
        for parcela in compra.parcelas:
            if inicio <= parcela.data_vencimento <= fim:
                conta["despesas"] += parcela.valor
                por_categoria[compra.categoria] = (
                    por_categoria.get(compra.categoria, Decimal("0")) + parcela.valor
                )
                mes = parcela.data_vencimento.strftime("%Y-%m")
                resumo_mes = por_mes.setdefault(
                    mes, {"receitas": Decimal("0"), "despesas": Decimal("0")}
                )
                resumo_mes["despesas"] += parcela.valor

    total_receitas = sum((receita.valor for receita in receitas), Decimal("0"))
    total_despesas = sum(
        (
            parcela.valor
            for compra in compras
            for parcela in compra.parcelas
            if inicio <= parcela.data_vencimento <= fim
        ),
        Decimal("0"),
    )
    return {
        "inicio": inicio,
        "fim": fim,
        "totais": {
            "receitas": total_receitas,
            "despesas": total_despesas,
            "saldo": total_receitas - total_despesas,
        },
        "por_conta": list(por_conta.values()),
        "por_categoria": [
            {"categoria": categoria, "valor": valor}
            for categoria, valor in sorted(
                por_categoria.items(), key=lambda item: item[1], reverse=True
            )
        ],
        "por_mes": [
            {"mes": mes, **valores}
            for mes, valores in sorted(por_mes.items())
        ],
    }


@router.get("/importacoes/modelo")
def baixar_modelo_importacao(
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
) -> StreamingResponse:
    conta = sessao.scalar(
        select(ContaFinanceira)
        .where(
            ContaFinanceira.usuario_id == usuario.id,
            ContaFinanceira.ativa.is_(True),
        )
        .order_by(ContaFinanceira.id)
    )
    nome_conta = conta.nome if conta is not None else "Nome da conta cadastrada"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Lançamentos"
    worksheet.append(list(IMPORT_HEADERS))
    worksheet.append(
        [
            "receita",
            "Salário",
            "Salário",
            5000,
            date.today(),
            nome_conta,
            1,
            "Exemplo de receita",
        ]
    )
    worksheet.append(
        [
            "compra",
            "Supermercado",
            "Alimentação",
            250.50,
            date.today(),
            nome_conta,
            1,
            "Exemplo de compra",
        ]
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = cell.font.copy(bold=True)
    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = 22
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)
    return StreamingResponse(
        arquivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo-lancamentos.xlsx"'},
    )


@router.post("/importacoes/lancamentos")
def importar_lancamentos(
    arquivo: UploadFile = File(...),
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    nome = (arquivo.filename or "").lower()
    if not nome.endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="Envie um arquivo XLSX.")
    conteudo = arquivo.file.read()
    if len(conteudo) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="O arquivo deve ter no máximo 10 MB.")
    try:
        workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except (TypeError, ValueError, OSError) as exc:
        raise HTTPException(status_code=422, detail="Arquivo XLSX inválido.") from exc

    worksheet = workbook.active
    linhas = worksheet.iter_rows(values_only=True)
    cabecalho = tuple(
        str(valor).strip().lower() if valor is not None else ""
        for valor in next(linhas, ())
    )
    if cabecalho != IMPORT_HEADERS:
        raise HTTPException(
            status_code=422,
            detail=f"Cabeçalho inválido. Use o modelo com as colunas: {', '.join(IMPORT_HEADERS)}.",
        )

    contas = {
        conta.nome.casefold(): conta
        for conta in sessao.scalars(
            select(ContaFinanceira).where(
                ContaFinanceira.usuario_id == usuario.id,
                ContaFinanceira.ativa.is_(True),
            )
        )
    }
    pendentes: list[tuple[str, dict[str, object]]] = []
    for numero_linha, valores in enumerate(linhas, start=2):
        if not any(valor is not None and str(valor).strip() for valor in valores):
            continue
        dados = dict(zip(IMPORT_HEADERS, valores, strict=False))
        tipo = _normalizar_tipo_importado(dados["tipo"], numero_linha)
        nome_conta = _texto_importado(dados["conta"], "conta", numero_linha)
        if nome_conta.casefold() not in contas:
            raise HTTPException(
                status_code=422,
                detail=f"Linha {numero_linha}: conta '{nome_conta}' não encontrada.",
            )
        try:
            valor = Decimal(str(dados["valor"])).quantize(Decimal("0.01"))
        except (ArithmeticError, ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: valor inválido.") from exc
        if valor <= 0:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: valor deve ser maior que zero.")
        quantidade = dados["quantidade_parcelas"] or 1
        try:
            quantidade = int(quantidade)
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: quantidade_parcelas inválida.") from exc
        if tipo == "receita" and quantidade != 1:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: receitas devem ter 1 parcela.")
        if not 1 <= quantidade <= 120:
            raise HTTPException(status_code=422, detail=f"Linha {numero_linha}: quantidade_parcelas deve estar entre 1 e 120.")
        pendentes.append(
            (
                tipo,
                {
                    "descricao": _texto_importado(dados["descricao"], "descricao", numero_linha),
                    "categoria": _texto_importado(dados["categoria"], "categoria", numero_linha),
                    "valor": valor,
                    "data": _ler_data_importada(dados["data"], numero_linha),
                    "conta_id": contas[nome_conta.casefold()].id,
                    "quantidade_parcelas": quantidade,
                    "observacoes": str(dados["observacoes"]).strip() if dados["observacoes"] is not None else None,
                },
            )
        )

    for tipo, dados in pendentes:
        if tipo == "receita":
            sessao.add(
                Receita(
                    usuario_id=usuario.id,
                    descricao=dados["descricao"],
                    categoria=dados["categoria"],
                    valor=dados["valor"],
                    data=dados["data"],
                    conta_id=dados["conta_id"],
                    observacoes=dados["observacoes"],
                )
            )
        else:
            compra = Compra(
                usuario_id=usuario.id,
                descricao=dados["descricao"],
                categoria=dados["categoria"],
                valor_total=dados["valor"],
                quantidade_parcelas=dados["quantidade_parcelas"],
                data_compra=dados["data"],
                conta_id=dados["conta_id"],
                observacoes=dados["observacoes"],
            )
            _criar_parcelas(compra)
            sessao.add(compra)
    sessao.commit()
    return {"importados": len(pendentes)}


@router.patch("/contas/{conta_id}", response_model=ContaFinanceiraResposta)
def atualizar_conta(
    conta_id: int,
    dados: ContaFinanceiraAtualizar,
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    conta = _validar_conta(conta_id, usuario.id, sessao)
    for campo, valor in dados.model_dump(exclude_unset=True).items():
        setattr(conta, campo, valor)
    sessao.commit()
    sessao.refresh(conta)
    return conta


@router.delete("/contas/{conta_id}", status_code=status.HTTP_204_NO_CONTENT)
def excluir_conta(
    conta_id: int,
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    conta = _validar_conta(conta_id, usuario.id, sessao)
    possui_receitas = sessao.scalar(
        select(Receita.id).where(Receita.conta_id == conta_id).limit(1)
    )
    possui_compras = sessao.scalar(
        select(Compra.id).where(Compra.conta_id == conta_id).limit(1)
    )
    if possui_receitas is not None or possui_compras is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Não é possível excluir uma conta que possui movimentações.",
        )
    sessao.delete(conta)
    sessao.commit()


@router.post("/receitas", response_model=ReceitaResposta, status_code=status.HTTP_201_CREATED)
def criar_receita(
    dados: ReceitaCriar,
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    _validar_conta(dados.conta_id, usuario.id, sessao)
    receita = Receita(**dados.model_dump(), usuario_id=usuario.id)
    sessao.add(receita)
    sessao.commit()
    sessao.refresh(receita)
    return receita


@router.get("/receitas", response_model=list[ReceitaResposta])
def listar_receitas(
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
    ano: int | None = None,
    mes: int | None = None,
    conta_id: int | None = None,
):
    if mes is not None and not 1 <= mes <= 12:
        raise HTTPException(status_code=422, detail="O mês deve estar entre 1 e 12.")
    filtros = [Receita.usuario_id == usuario.id]
    if ano is not None and mes is not None:
        inicio = date(ano, mes, 1)
        fim = date(ano + (mes == 12), 1 if mes == 12 else mes + 1, 1)
        filtros.extend([Receita.data >= inicio, Receita.data < fim])
    elif ano is not None:
        filtros.extend([Receita.data >= date(ano, 1, 1), Receita.data < date(ano + 1, 1, 1)])
    if conta_id is not None:
        filtros.append(Receita.conta_id == conta_id)
    return list(
        sessao.scalars(
            select(Receita)
            .where(*filtros)
            .order_by(Receita.data.desc())
        )
    )


@router.patch("/receitas/{receita_id}", response_model=ReceitaResposta)
def atualizar_receita(
        receita_id: int,
        dados: ReceitaAtualizar,
        usuario: Usuario = Depends(obter_usuario_atual),
        sessao: Session = Depends(obter_sessao),
):
        receita = sessao.scalar(
            select(Receita).where(
                Receita.id == receita_id, Receita.usuario_id == usuario.id
            )
        )
        if receita is None:
            raise HTTPException(status_code=404, detail="Receita não encontrada.")
        if dados.conta_id is not None:
            _validar_conta(dados.conta_id, usuario.id, sessao)
        for campo, valor in dados.model_dump(exclude_unset=True).items():
            setattr(receita, campo, valor)
        sessao.commit()
        sessao.refresh(receita)
        return receita


@router.delete("/receitas/{receita_id}", status_code=status.HTTP_204_NO_CONTENT)
def excluir_receita(
        receita_id: int,
        usuario: Usuario = Depends(obter_usuario_atual),
        sessao: Session = Depends(obter_sessao),
):
        receita = sessao.scalar(
            select(Receita).where(
                Receita.id == receita_id, Receita.usuario_id == usuario.id
            )
        )
        if receita is None:
            raise HTTPException(status_code=404, detail="Receita não encontrada.")
        sessao.delete(receita)
        sessao.commit()


@router.post("/compras", response_model=CompraResposta, status_code=status.HTTP_201_CREATED)
def criar_compra(
    dados: CompraCriar,
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
):
    _validar_conta(dados.conta_id, usuario.id, sessao)
    valor_parcela = (dados.valor_total / dados.quantidade_parcelas).quantize(
        Decimal("0.01"), rounding=ROUND_DOWN
    )
    valores = [valor_parcela] * dados.quantidade_parcelas
    valores[-1] += dados.valor_total - sum(valores)

    compra = Compra(**dados.model_dump(), usuario_id=usuario.id)
    compra.parcelas = [
        Parcela(
            numero=numero,
            valor=valor,
            data_vencimento=dados.data_compra + relativedelta(months=numero - 1),
        )
        for numero, valor in enumerate(valores, start=1)
    ]
    sessao.add(compra)
    sessao.commit()
    sessao.refresh(compra)
    return compra


@router.get("/compras", response_model=list[CompraResposta])
def listar_compras(
    usuario: Usuario = Depends(obter_usuario_atual),
    sessao: Session = Depends(obter_sessao),
    ano: int | None = None,
    mes: int | None = None,
    conta_id: int | None = None,
):
    if mes is not None and not 1 <= mes <= 12:
        raise HTTPException(status_code=422, detail="O mês deve estar entre 1 e 12.")
    filtros = [Compra.usuario_id == usuario.id]
    if conta_id is not None:
        filtros.append(Compra.conta_id == conta_id)
    if ano is not None and mes is not None:
        inicio = date(ano, mes, 1)
        fim = date(ano + (mes == 12), 1 if mes == 12 else mes + 1, 1)
        filtros.extend([Parcela.data_vencimento >= inicio, Parcela.data_vencimento < fim])
    elif ano is not None:
        filtros.extend(
            [
                Parcela.data_vencimento >= date(ano, 1, 1),
                Parcela.data_vencimento < date(ano + 1, 1, 1),
            ]
        )
    consulta = select(Compra).options(selectinload(Compra.parcelas))
    if ano is not None or mes is not None:
        consulta = consulta.join(Compra.parcelas)
    consulta = consulta.where(*filtros).order_by(Compra.data_compra.desc())
    resultado = sessao.scalars(consulta)
    return list(resultado.unique())


@router.patch("/compras/{compra_id}", response_model=CompraResposta)
def atualizar_compra(
        compra_id: int,
        dados: CompraAtualizar,
        usuario: Usuario = Depends(obter_usuario_atual),
        sessao: Session = Depends(obter_sessao),
):
        compra = sessao.scalar(
            select(Compra)
            .options(selectinload(Compra.parcelas))
            .where(Compra.id == compra_id, Compra.usuario_id == usuario.id)
        )
        if compra is None:
            raise HTTPException(status_code=404, detail="Compra não encontrada.")

        alteracoes = dados.model_dump(exclude_unset=True)
        if dados.conta_id is not None:
            _validar_conta(dados.conta_id, usuario.id, sessao)
        parcelas_pagas = any(parcela.status == "paga" for parcela in compra.parcelas)
        campos_estruturais = {"valor_total", "quantidade_parcelas", "data_compra"}
        if parcelas_pagas and campos_estruturais.intersection(alteracoes):
            raise HTTPException(
                status_code=409,
                detail="Não é possível alterar valores ou parcelas após um pagamento.",
            )

        for campo, valor in alteracoes.items():
            setattr(compra, campo, valor)

        if campos_estruturais.intersection(alteracoes):
            valor_parcela = (compra.valor_total / compra.quantidade_parcelas).quantize(
                Decimal("0.01"), rounding=ROUND_DOWN
            )
            valores = [valor_parcela] * compra.quantidade_parcelas
            valores[-1] += compra.valor_total - sum(valores)
            compra.parcelas.clear()
            sessao.flush()
            compra.parcelas = [
                Parcela(
                    numero=numero,
                    valor=valor,
                    data_vencimento=compra.data_compra + relativedelta(months=numero - 1),
                )
                for numero, valor in enumerate(valores, start=1)
            ]

        sessao.commit()
        sessao.refresh(compra)
        return compra


@router.delete("/compras/{compra_id}", status_code=status.HTTP_204_NO_CONTENT)
def excluir_compra(
        compra_id: int,
        usuario: Usuario = Depends(obter_usuario_atual),
        sessao: Session = Depends(obter_sessao),
):
        compra = sessao.scalar(
            select(Compra).where(
                Compra.id == compra_id, Compra.usuario_id == usuario.id
            )
        )
        if compra is None:
            raise HTTPException(status_code=404, detail="Compra não encontrada.")
        sessao.delete(compra)
        sessao.commit()
